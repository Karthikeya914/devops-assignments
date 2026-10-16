import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    base_dir = "/Users/karthikeya/devops-assignments/session21"
    screenshots_dir = os.path.join(base_dir, "screenshots")
    output_pptx = os.path.join(base_dir, "demo", "presentation.pptx")

    # Colors
    BG_COLOR = RGBColor(15, 23, 42)      # Slate 900
    CARD_BG = RGBColor(30, 41, 59)      # Slate 800
    TEXT_MAIN = RGBColor(248, 250, 252) # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)# Slate 400
    ACCENT_CYAN = RGBColor(6, 182, 212) # Cyan 500
    ACCENT_BLUE = RGBColor(59, 130, 246) # Blue 500
    ACCENT_GREEN = RGBColor(34, 197, 94)# Green 500

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, title_text, category_text="OPSPULSE DEVOPS CAPSTONE"):
        # Category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

    # --- SLIDE 1: Title Slide ---
    slide_layout = prs.slide_layouts[6] # Blank
    s1 = prs.slides.add_slide(slide_layout)
    set_slide_background(s1)

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p0 = tf1.paragraphs[0]
    p0.text = "OPSPULSE PLATFORM"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN

    p1 = tf1.add_paragraph()
    p1.text = "Cloud Incident & Reliability Engineering Platform"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_MAIN

    p2 = tf1.add_paragraph()
    p2.text = "DevOps Final Capstone Project (Session 21) • 100-Point Architecture Walkthrough"
    p2.font.size = Pt(18)
    p2.font.color.rgb = TEXT_MUTED

    p3 = tf1.add_paragraph()
    p3.text = "\nPresenter: Karthikeya   |   Repository: github.com/Karthikeya914/devops-assignments"
    p3.font.size = Pt(14)
    p3.font.color.rgb = ACCENT_BLUE

    # Helper function for Content + Screenshot slide
    def make_content_image_slide(title, category, bullets, image_filename, img_left=6.8, img_top=1.6, img_w=5.7):
        s = prs.slides.add_slide(slide_layout)
        set_slide_background(s)
        add_header(s, title, category)

        # Bullets box
        box = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
        tf = box.text_frame
        tf.word_wrap = True

        for i, (b_title, b_desc) in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {b_title}: "
            p.font.bold = True
            p.font.size = Pt(14)
            p.font.color.rgb = ACCENT_CYAN
            
            run = p.add_run()
            run.text = b_desc
            run.font.bold = False
            run.font.size = Pt(13)
            run.font.color.rgb = TEXT_MAIN
            p.space_after = Pt(14)

        # Image
        img_path = os.path.join(screenshots_dir, image_filename)
        if os.path.exists(img_path):
            s.shapes.add_picture(img_path, Inches(img_left), Inches(img_top), width=Inches(img_w))
        return s

    # --- SLIDE 2: Application Stack (M1) ---
    make_content_image_slide(
        title="Module 1: Full-Stack Polyglot Application",
        category="M1 — APPLICATION ARCHITECTURE",
        bullets=[
            ("FastAPI Backend", "High-performance Python asynchronous REST API exposing /health, /ready, and incident CRUD endpoints."),
            ("React + Vite Frontend", "Interactive single-page application with real-time KPI metrics, incident status filters, and dark mode UI."),
            ("PostgreSQL Database", "Managed relational storage with schema migrations version-controlled via Alembic."),
            ("Local Orchestration", "Single-command bootstrapping with Docker Compose connecting backend, frontend, and database.")
        ],
        image_filename="01_app_dashboard_live.png"
    )

    # --- SLIDE 3: Automated Testing & Quality Gates (M2) ---
    make_content_image_slide(
        title="Module 2: Automated Testing & Quality Gates",
        category="M2 — AUTOMATED TESTING",
        bullets=[
            ("Pytest Suite", "10 comprehensive automated test cases covering health probes, root info, CRUD operations, and error handling."),
            ("Isolated Test Database", "In-memory SQLite fixtures ensure tests never touch or mutate production database state."),
            ("Quality Gate Enforcement", "100% test pass rate executed in 0.12s before code can be packaged or deployed."),
            ("Clean Configuration", "Standardized pytest.ini and conftest.py fixtures for reproducible execution.")
        ],
        image_filename="03_pytest_automated_tests.png"
    )

    # --- SLIDE 4: Docker Containerization (M4) ---
    make_content_image_slide(
        title="Module 4: Multi-Stage Docker Containerization",
        category="M4 — CONTAINERIZATION",
        bullets=[
            ("Multi-Stage Builds", "Frontend uses Node 20 builder stripped down to an ultra-lightweight Nginx Alpine runtime image."),
            ("Non-Root Security", "Both backend and frontend containers run under unprivileged user 'appuser' (UID 10001)."),
            ("Optimized Image Sizes", "Backend is 196MB and Frontend is 28MB, minimizing attack surface and pull latency."),
            ("Docker Compose", "Orchestrates multi-container networking with internal DNS resolution and persistent volumes.")
        ],
        image_filename="04_docker_compose_and_images.png"
    )

    # --- SLIDE 5: CI/CD Pipeline & GHCR (M5) ---
    make_content_image_slide(
        title="Module 5: Automated CI/CD with GitHub Actions",
        category="M5 — CONTINUOUS INTEGRATION",
        bullets=[
            ("Triggers on Push", "Automated pipeline triggers on every commit to main branch preventing untested code deployment."),
            ("Automated Pipeline Stages", "Job 1 runs Pytest suite. Job 2 compiles frontend and builds multi-stage Docker images."),
            ("GitHub Container Registry", "Images published to ghcr.io/karthikeya914 tagged immutably with Git commit SHA."),
            ("Zero Manual Intervention", "100% hands-free release packaging with automated secrets management.")
        ],
        image_filename="06_github_actions_pipeline_runs.png"
    )

    # --- SLIDE 6: DevSecOps Vulnerability Scanning (M6) ---
    make_content_image_slide(
        title="Module 6: DevSecOps & Security Hardening",
        category="M6 — SECURITY GATING",
        bullets=[
            ("Automated Trivy Scanning", "Scans both backend and frontend container images for known CVE vulnerabilities in CI."),
            ("Severity Gating", "Configured to fail pipelines on CRITICAL or HIGH severity security flaws."),
            ("Zero-Vulnerability Baseline", "Multi-stage alpine base images eliminate package bloat, achieving 0 vulnerabilities."),
            ("Secret Defense", "Rigorous .gitignore hygiene ensures no API keys, credentials, or .env files enter version control.")
        ],
        image_filename="09_trivy_security_scan_clean.png"
    )

    # --- SLIDE 7: Terraform Infrastructure as Code (M7) ---
    make_content_image_slide(
        title="Module 7: Infrastructure as Code on AWS",
        category="M7 — CLOUD INFRASTRUCTURE",
        bullets=[
            ("Modular Terraform HCL", "Declares dedicated AWS VPC (10.0.0.0/16) and managed EKS cluster in ap-south-1 (Mumbai)."),
            ("High Availability", "Spans 2 Availability Zones with public subnets, private subnets, and NAT Gateway egress."),
            ("Live AWS Verification", "EKS cluster 'opspulse-eks' and VPC 'opspulse-eks-vpc' verified active in AWS Console."),
            ("Clean Teardown", "Automated 'terraform destroy' confirmed all 57 cloud resources cleanly deleted with zero runaway costs.")
        ],
        image_filename="14_aws_console_eks_cluster.png"
    )

    # --- SLIDE 8: Kubernetes Orchestration & Helm (M8) ---
    make_content_image_slide(
        title="Module 8: Kubernetes & Helm Packaging",
        category="M8 — ORCHESTRATION & PACKAGING",
        bullets=[
            ("Helm Chart Architecture", "Parameterized chart in helm/opspulse managing Deployments, Services, Ingress, and HPA."),
            ("High Availability", "Dual-replica backend and frontend deployments guaranteeing zero downtime during rolling updates."),
            ("Health Probes", "Configured livenessProbe (/health) and readinessProbe (/ready) for automatic pod self-healing."),
            ("Ingress Routing", "Nginx Ingress controller routes external traffic to frontend (/) and API requests to backend (/api).")
        ],
        image_filename="11_k8s_running_pods.png"
    )

    # --- SLIDE 9: Observability & Monitoring (M9) ---
    make_content_image_slide(
        title="Module 9: Observability & Telemetry",
        category="M9 — MONITORING & TELEMETRY",
        bullets=[
            ("Prometheus Telemetry", "Backend exports real-time HTTP metrics, status counters, and incident telemetry on /metrics."),
            ("Prometheus Scraping", "Configured scrape targets and Helm values for automated metrics collection."),
            ("Grafana Visualizations", "Live dashboards monitoring request rates, p95 latency, error spikes, and active incident volume."),
            ("Proactive Alerting", "Preconfigured alert rules for high error rates and database pool exhaustion.")
        ],
        image_filename="10_prometheus_metrics_endpoint.png"
    )

    # --- SLIDE 10: Key Takeaways & Summary ---
    s10 = prs.slides.add_slide(slide_layout)
    set_slide_background(s10)
    add_header(s10, "Module 10: Key Learnings & Engineering Summary", "M10 — RETROSPECTIVE")

    box10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(5.2))
    tf10 = box10.text_frame
    tf10.word_wrap = True

    takeaways = [
        ("1. Shift-Left Security Works", "Integrating Trivy container scans into GitHub Actions eliminates CVEs before deployment."),
        ("2. Infrastructure as Code is Reproducible", "Terraform allows provisioning and tearing down a complete AWS VPC & EKS environment in minutes."),
        ("3. Observability vs Monitoring", "Exposing rich Prometheus metrics and readiness probes enables Kubernetes to make intelligent routing decisions and self-heal failed pods."),
        ("4. Immutable GitOps Releases", "Tagging images with Git commit SHAs and reconciling cluster state via Helm charts creates a fully auditable production deployment history.")
    ]

    for i, (heading, body) in enumerate(takeaways):
        p = tf10.paragraphs[0] if i == 0 else tf10.add_paragraph()
        p.text = f"{heading}: "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = ACCENT_CYAN

        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.size = Pt(14)
        run.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(16)

    prs.save(output_pptx)
    print(f"Presentation successfully created at: {output_pptx}")

if __name__ == "__main__":
    create_presentation()
